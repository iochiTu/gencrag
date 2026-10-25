"""Excel/CSV import: reading, validation, description reading, and the full admin flow."""
import base64
import io
import os
import time

import pytest

from app import importer
from app.importer import extract_attributes, ingredient_flags, parse_ingredients, read_table


def xlsx(rows, headers=("SKU", "Name", "Price", "Description")):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_template_round_trips():
    p = read_table("template.xlsx", importer.template_xlsx())
    assert len(p["rows"]) == 2 and all(r["ok"] for r in p["rows"]) and not p["issues"]
    assert p["rows"][0]["category"] == "serum" and p["rows"][0]["cost"] == 7.2
    assert [r["margin_pct"] for r in p["rows"]] == [0.70, 0.78]
    assert [r["priority"] for r in p["rows"]] == [5, 3]


def test_excel_percentage_cells_and_plain_values():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["SKU", "Margin (%)", "Manual Priority"])
    for row in [("A", 0.4, 5), ("B", 40, 3), ("C", "40%", 1), ("D", 0, ""), ("E", "100%", 4)]:
        ws.append(row)
    ws["B2"].number_format = "0.0%"
    buf = io.BytesIO()
    wb.save(buf)
    p = read_table("margins.xlsx", buf.getvalue())
    assert not p["issues"]
    assert [r["margin_pct"] for r in p["rows"]] == [0.4, 0.4, 0.4, 0, 1]
    assert [r["priority"] for r in p["rows"]] == [5, 3, 1, None, 4]


@pytest.mark.parametrize("margin,priority", [("-1%", "0"), ("101%", "6"), ("NaN", "2.5"), ("inf", "inf"), ("abc", "popular")])
def test_bad_margin_and_priority_warn_without_overwriting(margin, priority):
    p = read_table("business.csv", f"SKU,Margin,Priority\nA,{margin},{priority}\n".encode())
    assert p["rows"][0]["ok"]
    assert p["rows"][0]["margin_pct"] is None and p["rows"][0]["priority"] is None
    assert len(p["issues"]) == 2 and all(i["level"] == "warning" for i in p["issues"])


def test_header_synonyms_prices_and_numeric_skus():
    data = xlsx([(1001, "Toner", "$18.50", "x" * 80 + " Ingredients: Water, Glycerin"),
                 ("A-2", "Cream", "1,299.00", "y" * 80 + " Ingredients: Water")],
                headers=("Item Code", "Product Name", "Retail Price", "Details"))
    rows = read_table("f.xlsx", data)["rows"]
    assert [r["sku"] for r in rows] == ["1001", "A-2"] and [r["price"] for r in rows] == [18.5, 1299.0]


class FakeConn:
    """Stands in for the database in finalise_rows: knows which SKUs already exist."""
    def __init__(self, skus): self.skus = skus
    def execute(self, *_):
        class R:
            def __init__(s, rows): s.rows = rows
            def fetchall(s): return s.rows
        return R([{"sku": x} for x in self.skus])


def test_validation_messages_use_excel_row_numbers():
    data = xlsx([("S1", "A", 10, "desc " * 20), ("S1", "B", 12, "desc " * 20), ("", "C", 5, "d"),
                 ("S3", "D", "free", "desc " * 20), ("S4", "", 5, "")])
    p = read_table("f.xlsx", data)
    importer.finalise_rows(FakeConn([]), p)
    errs = {(i["row"], i["message"].split(" (")[0].split(";")[0]) for i in p["issues"] if i["level"] == "error"}
    assert (3, "SKU S1 appears twice") in errs and (4, "SKU is empty") in errs
    assert (5, "New product needs Price") in errs and (6, "New product needs Name") in errs
    warns = [i["message"] for i in p["issues"] if i["level"] == "warning"]
    assert any("'free' is not a number" in w for w in warns)
    assert [r["ok"] for r in p["rows"]] == [True, False, False, False, False]


def test_messy_cells_are_tolerated():
    from datetime import date, datetime
    data = xlsx([
        ("A1", "Toner", "N/A", "#N/A", "-", "TBD", "31/12/2027", ""),             # existing: blanks keep values
        ("A2", "Serum", "$24.00", "Serum for dry skin. " * 4, "abc", "lots", "2027-06", "Water, Retinol, Parfum"),
        ("A3", "Cream", 30, None, None, 12, datetime(2027, 3, 4), None),           # Excel date cell
        ("A4", "Mask", 3, "Sheet mask for dull skin, all skin types.", "", "", "03/04/2027", "  "),
        ("A5", "Old", 9, "Old cream for dry skin and fine lines.", "", "", "2020-01-01", ""),
        (None, None, "Total", None, None, None, None, None),                        # totals row: ignored
    ], headers=("SKU", "Name", "Price", "Description", "Cost", "Stock", "Expiry", "Ingredients"))
    p = read_table("f.xlsx", data, date_order="MDY")
    importer.finalise_rows(FakeConn(["A1"]), p)
    rows = {r["sku"]: r for r in p["rows"]}
    assert set(rows) == {"A1", "A2", "A3", "A4", "A5"} and all(r["ok"] for r in rows.values())
    a1 = rows["A1"]
    assert a1["price"] is None and a1["description"] is None and a1["cost"] is None and a1["stock"] is None
    assert a1["expiry"] == "2027-12-31"                                            # day > 12: unambiguous
    assert rows["A2"]["price"] == 24 and rows["A2"]["cost"] is None and rows["A2"]["expiry"] == "2027-06-30"
    assert rows["A3"]["expiry"] == "2027-03-04" and rows["A4"]["expiry"] == "2027-03-04"
    msgs = " | ".join(i["message"] for i in p["issues"])
    assert "Empty price, description: current value kept" in msgs                   # A1 exists
    assert "Cost 'abc' is not a number" in msgs and "Stock 'lots'" in msgs
    assert "read as 04 Mar 2027" in msgs                                            # ambiguous date explained
    assert "No description" in msgs                                                 # A3 is new, no description
    assert "is in the past" in msgs                                                 # A5 expired
    d, note = importer.parse_date("03/04/2027", "DMY")
    assert d == date(2027, 4, 3) and "day/month" in note
    assert importer.parse_date("Mar 2027")[0] == date(2027, 3, 31)
    assert importer.parse_date("20270315")[0] == date(2027, 3, 15)
    assert importer.parse_date("soon")[0] is None


def test_ingredients_column_is_used_when_description_has_none():
    assert importer.full_description("Night serum.", "Water, Retinol") == "Night serum.\nIngredients: Water, Retinol"
    assert importer.full_description("Serum. Ingredients: Water", "Water, Retinol") == "Serum. Ingredients: Water"
    assert importer.full_description(None, "Water") == "Ingredients: Water"


def test_update_only_file_with_sku_and_price():
    p = read_table("prices.csv", b"SKU,Price\nA1,19.99\nA2,\n")
    importer.finalise_rows(FakeConn(["A1", "A2"]), p)
    assert [r["ok"] for r in p["rows"]] == [True, True] and p["rows"][1]["price"] is None
    with pytest.raises(ValueError, match="Only a SKU column"):
        read_table("x.csv", b"SKU\nA1\n")


def test_flags_from_description_text_when_no_ingredient_list():
    a = extract_attributes(None, "Night Serum", "Night serum with 0.3% retinol for fine lines. Fragrance-free.", use_ai=False)
    assert a["text_flags"] == ["retinoid"] and "pregnancy" in a["avoid_for"] and a["needs_review"]
    assert "description mentions: retinoid" in a["review_note"]
    b = extract_attributes(None, "Mystery", "", use_ai=False)
    assert b["needs_review"] and "no description" in b["review_note"]


def test_missing_required_column_is_explained():
    with pytest.raises(ValueError, match="Missing column: SKU"):
        read_table("f.xlsx", xlsx([("A", 10, "d")], headers=("Name", "Price", "Description")))
    with pytest.raises(ValueError, match="Excel"):
        read_table("f.pdf", b"%PDF")


def test_ingredients_and_flags():
    ings = parse_ingredients("Gentle toner.\nIngredients: Water, Glycolic Acid (5%), Parfum, "
                             "Cetearyl Alcohol, Alcohol Denat., Melaleuca Alternifolia (Tea Tree) Leaf Oil.")
    assert ings[:3] == ["Water", "Glycolic Acid (5%)", "Parfum"]
    flags = {i: ingredient_flags(i) for i in ings}
    assert flags["Glycolic Acid (5%)"] == ["aha"] and flags["Parfum"] == ["fragrance"]
    assert flags["Cetearyl Alcohol"] == []                                      # fatty alcohol is fine
    assert flags["Alcohol Denat"] == ["alcohol"]                                # trailing period stripped
    assert "essential_oil" in flags["Melaleuca Alternifolia (Tea Tree) Leaf Oil"]


class FakeLLM:
    name = "fake"
    def __init__(self, reply): self.reply = reply
    def chat(self, system, user, role="rank"): return self.reply
    def model_label(self, role): return "fake"


def test_ai_reading_is_merged_safety_first():
    desc = ("Overnight renewal cream for all skin types, ages 30-55, softens fine lines. "
            "Avoid during pregnancy.\nIngredients: Water, Retinol, Squalane, Fragrance")
    ai = FakeLLM('{"category":"cream","suitable_skin_types":["dry","normal","sensitive"],"target_concerns":["wrinkles"],'
                 '"avoid_for":[],"min_age":30,"max_age":55,"key_benefits":["Softens fine lines overnight"],"ingredients":[]}')
    a = extract_attributes(ai, "Renewal Night Cream", desc)
    assert a["category"] == "cream" and a["extracted_by"] == "ai+rules"
    assert "pregnancy" in a["avoid_for"]                                        # AI missed it; rules + retinoid keep it
    assert "sensitive" not in a["suitable_skin_types"]                         # strong actives, not stated as sensitive-safe
    assert set(a["flags"]) == {"retinoid", "fragrance"} and a["fragrance_free"] is False
    assert (a["min_age"], a["max_age"]) == (30, 55) and a["key_benefits"] == ["Softens fine lines overnight"]
    assert a["ingredients"] == ["Water", "Retinol", "Squalane", "Fragrance"]   # list taken from the text, not the AI


def test_rules_only_when_ai_fails_and_review_flag():
    a = extract_attributes(FakeLLM("sorry"), "Mystery Product", "Nice product for everyone.")
    assert a["extracted_by"] == "rules" and a["needs_review"]
    assert "AI could not read" in a["review_note"] and "no ingredient list" in a["review_note"]
    b = extract_attributes(None, "Calming Toner", "Toner that calms redness. Not suitable for sensitive skin. "
                           "For oily and combination skin.\nIngredients: Water, Salicylic Acid", use_ai=False)
    assert b["category"] == "toner" and "sensitive" not in b["suitable_skin_types"]
    assert {"oily", "combination"} <= set(b["suitable_skin_types"]) and b["flags"] == ["bha"]


# ------------------------------------------------------------------ end-to-end through the admin API
pytestmark_db = pytest.mark.skipif(not os.environ.get("TEST_DATABASE_URL"), reason="TEST_DATABASE_URL not set")


@pytest.fixture(scope="module")
def client():
    if not os.environ.get("TEST_DATABASE_URL"):
        pytest.skip("TEST_DATABASE_URL not set")
    from fastapi.testclient import TestClient
    from app import db
    from app.main import app
    db.run_sql_file("schema.sql")
    db.run_sql_file("seed_sample.sql")
    c = TestClient(app)
    assert c.post("/admin/api/login", json={"name": "Ops", "password": "test-pass"}).status_code == 200
    from tests.conftest import staff_login
    staff_login(c)
    return c


def upload_and_run(client, data, filename="products.xlsx", **opts):
    pv = client.post("/admin/api/import/preview", json={"filename": filename,
                                                        "content_b64": base64.b64encode(data).decode()})
    assert pv.status_code == 200, pv.text
    pv = pv.json()
    r = client.post(f"/admin/api/import/{pv['job_id']}/run", json=opts)
    assert r.status_code == 200, r.text
    for _ in range(100):
        job = client.get(f"/admin/api/import/{pv['job_id']}").json()
        if job["status"] in ("done", "failed"):
            return pv, job
        time.sleep(0.1)
    raise AssertionError("import did not finish")


NEW = [("NW-001", "Cica Calming Serum", 24,
        "Lightweight serum that calms redness and helps with breakouts. For sensitive, oily and combination skin. "
        "Fragrance-free.\nIngredients: Water, Centella Asiatica Extract, Niacinamide, Panthenol"),
       ("NW-002", "Retinol Night Cream", 48,
        "Night cream for dry, mature skin (40+) that softens fine lines. Not recommended during pregnancy.\n"
        "Ingredients: Water, Retinol, Shea Butter, Fragrance"),
       ("NW-003", "Oil-Free Cica Gel Serum", 19,
        "Oil-free gel serum for oily and combination skin with breakouts and redness. Fragrance-free.\n"
        "Ingredients: Water, Centella Asiatica Extract, Zinc PCA, Allantoin")]


@pytestmark_db
def test_import_preview_run_and_recommend(client):
    tmpl = client.get("/admin/api/import/template")
    assert tmpl.status_code == 200 and tmpl.content[:2] == b"PK"
    pv, job = upload_and_run(client, xlsx(NEW))
    assert pv["summary"]["new"] == 3 and pv["summary"]["missing"] == 15        # the 15 seed products not in file
    assert job["status"] == "done", job
    c = job["summary"]["counts"]
    assert c["created"] == 3 and c["analyzed"] == 3 and c["embedded"] >= 3
    res = {r["sku"]: r for r in job["results"]}
    assert res["NW-002"]["category"] == "cream" and "pregnancy" in res["NW-002"]["avoid"]
    detail = client.get("/admin/api/products/NW-002").json()["product"]
    assert [i["inci_name"] for i in detail["ingredients"]][:2] == ["Water", "Retinol"]

    r = client.post("/recommend", json={"profile": {"skin_type": "oily", "sensitive": True,
                                                    "concerns": ["redness", "acne"]}, "top_n": 3, "options": 3}).json()
    skus = [o["sku"] for st in r["steps"] for o in st["options"]]
    assert "NW-001" in skus or "NW-003" in skus
    serum = next(st for st in r["steps"] if st["category"] == "serum")
    assert len(serum["options"]) >= 2                                          # staff get a choice
    preg = client.post("/recommend", json={"text": "dry mature skin, 45, fine lines, pregnant", "top_n": 4}).json()
    assert "NW-002" not in [o["sku"] for st in preg.get("steps", []) for o in st["options"]]


@pytestmark_db
def test_reimport_only_reanalyses_changed_rows_and_can_deactivate(client):
    rows = [list(r) for r in NEW]
    rows[0][2] = 26                                                            # price change only
    rows[2][3] = rows[2][3].replace("breakouts", "breakouts and enlarged pores")   # description change
    pv, job = upload_and_run(client, xlsx(rows), deactivate_missing=True)
    assert pv["summary"]["changed"] == 2 and pv["summary"]["unchanged"] == 1
    c = job["summary"]["counts"]
    assert c["analyzed"] == 1 and c["updated"] == 2 and c["unchanged"] == 1 and c["deactivated"] == 15
    prods = {p["sku"]: p for p in client.get("/admin/api/products").json()["products"]}
    assert prods["NW-001"]["price"] == 26 and prods["CL-001"]["is_active"] is False
    again = client.post(f"/admin/api/import/{pv['job_id']}/run", json={})
    assert again.status_code == 409


@pytestmark_db
def test_manual_edit_of_ai_reading(client):
    body = {"category": "serum", "suitable_skin_types": ["oily", "combination"], "target_concerns": ["acne", "pores"],
            "avoid_for": [], "min_age": 16, "max_age": None, "key_benefits": ["Clears pores", ""], "is_active": True}
    r = client.put("/admin/api/products/NW-003/attributes", json=body)
    assert r.status_code == 200 and r.json()["warning"] is None
    p = client.get("/admin/api/products/NW-003").json()["product"]
    assert p["extracted_by"] == "manual" and p["needs_review"] is False and p["key_benefits"] == ["Clears pores"]
    assert p["searchable"] is True
    bad = client.put("/admin/api/products/NW-003/attributes", json={**body, "min_age": 50, "max_age": 20})
    assert bad.status_code == 422


@pytestmark_db
def test_preview_rejects_bad_files(client):
    r = client.post("/admin/api/import/preview", json={"filename": "x.xlsx", "content_b64": base64.b64encode(b"nope").decode()})
    assert r.status_code == 400 and "Could not read" in r.json()["detail"]


@pytestmark_db
def test_import_with_blanks_keeps_existing_values_and_hides_unsafe(client):
    upload_and_run(client, xlsx(NEW))
    rows = [("NW-001", "", "", "", "2027-12-31"),                                   # blanks: keep, add expiry
            ("NW-004", "Mystery Night Serum", 33, "Night serum with retinol for fine lines, dry skin.", "N/A"),
            ("NW-005", "Expired Toner", 12, "Toner for oily skin with breakouts.\nIngredients: Water", "2020-01-01")]
    pv, job = upload_and_run(client, xlsx(rows, headers=("SKU", "Name", "Price", "Description", "Expiry")))
    assert job["status"] == "done", job
    prods = {p["sku"]: p for p in client.get("/admin/api/products").json()["products"]}
    assert prods["NW-001"]["name"] == "Cica Calming Serum" and prods["NW-001"]["price"] == 24
    assert prods["NW-001"]["expiry_date"] == "2027-12-31"
    assert prods["NW-004"]["ingredients_known"] is False and prods["NW-004"]["needs_review"] is True
    # retinol mentioned in text -> hidden for pregnancy even without an ingredient list
    r = client.post("/recommend", json={"profile": {"concerns": ["wrinkles"], "avoid": ["pregnancy"]}, "top_n": 5}).json()
    assert "NW-004" not in [o["sku"] for st in r.get("steps", []) for o in st["options"]]
    # unknown ingredients: hidden when the customer must avoid fragrance, shown with a warning if set to "show"
    q = {"profile": {"concerns": ["wrinkles"], "skin_type": "dry", "avoid": ["fragrance"]}, "top_n": 5, "options": 4}
    assert "NW-004" not in [o["sku"] for st in client.post("/recommend", json=q).json().get("steps", []) for o in st["options"]]
    rk = client.get("/admin/api/ranking").json()["settings"]
    client.put("/admin/api/ranking", json={**rk, "unknown_ingredients": "show"})
    shown = [o for st in client.post("/recommend", json=q).json()["steps"] for o in st["options"] if o["sku"] == "NW-004"]
    client.put("/admin/api/ranking", json=rk)
    assert shown and any("check the label" in n for n in shown[0]["notes"])
    # expired products are never recommended
    r = client.post("/recommend", json={"profile": {"skin_type": "oily", "concerns": ["acne"]}, "top_n": 5, "options": 4}).json()
    assert "NW-005" not in [o["sku"] for st in r["steps"] for o in st["options"]]


@pytestmark_db
def test_margin_and_priority_import_updates_ranking_without_reanalysis(client):
    from app.db import connect
    from app.scoring import score_candidates
    from app.models import CustomerProfile
    desc = "Gentle serum for dry skin and dehydration. Fragrance-free. Ingredients: Water, Glycerin, Panthenol"
    rows = [("MP-A", "Hydrating Serum", 10, desc, "40%", 1),
            ("MP-B", "Hydrating Serum", 30, desc, "40%", 5)]
    pv, job = upload_and_run(client, xlsx(rows, headers=("SKU", "Name", "Price", "Description", "Margin (%)", "Priority")))
    assert job["status"] == "done" and job["summary"]["counts"]["created"] == 2
    prods = {p["sku"]: p for p in client.get("/admin/api/products").json()["products"]}
    assert prods["MP-B"]["margin_pct"] == 0.4 and prods["MP-B"]["priority"] == 5

    # Identical customer fit and margin: manually favored product is first, despite its higher price.
    r = client.post("/recommend", json={"profile": {"categories": ["serum"], "concerns": ["dryness"]}, "options": 4}).json()
    assert r["steps"][0]["options"][0]["sku"] == "MP-B"
    budget = client.post("/recommend", json={"profile": {"categories": ["serum"], "budget_max": 20}, "options": 4}).json()
    assert "MP-B" not in [o["sku"] for st in budget["steps"] for o in st["options"]]

    with connect() as conn:
        before = conn.execute("SELECT embedding::text AS embedding FROM products WHERE sku = 'MP-A'").fetchone()
        conn.execute("UPDATE business_rules SET campaign_boost = 0.07, excluded = TRUE WHERE sku = 'MP-A'")
    pv, job = upload_and_run(client, b"SKU,Margin,Priority\nMP-A,80%,5\n", filename="business.csv")
    counts = job["summary"]["counts"]
    assert pv["summary"]["changed"] == 1
    assert "margin → 80%" in pv["diff"]["changed"][0]["changes"]
    assert "priority 1 → 5" in pv["diff"]["changed"][0]["changes"]
    assert counts["updated"] == 1 and counts["analyzed"] == 0 and counts["embedded"] == 0
    with connect() as conn:
        after = conn.execute("SELECT margin_pct::float AS margin, embedding::text AS embedding FROM products WHERE sku = 'MP-A'").fetchone()
        rule = conn.execute("SELECT * FROM business_rules WHERE sku = 'MP-A'").fetchone()
    assert after["margin"] == 0.8 and after["embedding"] == before["embedding"]
    assert rule["priority"] == 5 and rule["excluded"] and float(rule["campaign_boost"]) == 0.07 and rule["updated_by"] == "Ops"

    # Blank business cells and price/cost updates keep the direct margin and manual priority.
    _, job = upload_and_run(client, b"SKU,Price,Cost,Margin,Priority\nMP-A,50,49,,\n", filename="prices.csv")
    assert job["summary"]["counts"]["analyzed"] == 0
    p = next(p for p in client.get("/admin/api/products").json()["products"] if p["sku"] == "MP-A")
    assert p["margin_pct"] == 0.8 and p["priority"] == 5 and p["price"] == 50
    assert client.get("/admin/api/products/MP-A").json()["product"]["margin_pct"] == 0.8
    assert client.get("/admin/api/analytics").status_code == 200
    scored = score_candidates([{**p, "similarity": 0.5}, {**p, "price": 1, "similarity": 0.5}], CustomerProfile())
    assert scored[0]["score"] == scored[1]["score"]
    _, job = upload_and_run(client, b"SKU,Margin,Priority\nMP-A,0%,3\n", filename="zero.csv")
    assert job["summary"]["counts"]["updated"] == 1
    p = next(p for p in client.get("/admin/api/products").json()["products"] if p["sku"] == "MP-A")
    assert p["margin_pct"] == 0 and p["priority"] == 3
